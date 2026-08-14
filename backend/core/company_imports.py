import re
from io import BytesIO

from fastapi import HTTPException
from openpyxl import load_workbook

from .companies import CompanyCreate, create_company_for_admin
from .mandates import MandateCreate, create_mandate_for_admin
from .permissions import require_tenant_context

ALIASES={"name":["name","nom","société","societe"],"legal_name":["legal name","nom légal","nom legal","raison sociale"],"company_code":["company code","code société","code societe","code"],"jurisdiction":["jurisdiction","juridiction","pays"],"region":["region","région","canton","province"],"functional_currency":["functional currency","devise","devise fonctionnelle","currency"],"industry":["industry","secteur"],"company_type":["company type","type société","type societe","type"],"fiscal_year_start":["fiscal year start","début exercice","debut exercice"],"mandate_code":["mandate code","code mandat","mandat"],"principal_email":["principal email","responsable","responsable principal","email responsable"],"collaborator_emails":["collaborators","collaborateurs","collaborateurs autorisés","collaborateurs autorises"]}

def norm(v): return re.sub(r"\s+"," ",str(v or "").strip().lower().replace("_"," "))
def split_emails(v): return list(dict.fromkeys(x.strip().lower() for x in re.split(r"[;,\n]+",str(v or "")) if x.strip()))
def parse_workbook(content,fiduciary=False):
    try: ws=load_workbook(BytesIO(content),read_only=True,data_only=True).active
    except Exception as e: raise HTTPException(422,"Fichier Excel invalide") from e
    rows=list(ws.iter_rows(values_only=True))
    if not rows: raise HTTPException(422,"Le fichier Excel est vide")
    alias={a:k for k,vals in ALIASES.items() for a in vals}; cols={alias[norm(v)]:i for i,v in enumerate(rows[0]) if norm(v) in alias}
    if "name" not in cols: raise HTTPException(422,"Colonne obligatoire manquante: Nom")
    def cell(row,k,default=None):
        i=cols.get(k); v=row[i] if i is not None and i<len(row) else None
        return str(v).strip() if v is not None else default
    out=[]
    for rn,row in enumerate(rows[1:],2):
        if not any(v not in (None,"") for v in row): continue
        j=(cell(row,"jurisdiction","") or "").upper(); d={"row":rn,"name":cell(row,"name","") or "","legal_name":cell(row,"legal_name"),"company_code":cell(row,"company_code"),"jurisdiction":j,"region":cell(row,"region"),"functional_currency":(cell(row,"functional_currency",{"CH":"CHF","CA":"CAD"}.get(j,"")) or "").upper(),"industry":cell(row,"industry","services") or "services","company_type":cell(row,"company_type","operating") or "operating","fiscal_year_start":cell(row,"fiscal_year_start","01-01") or "01-01"}
        if fiduciary: d.update(mandate_code=cell(row,"mandate_code"),principal_email=(cell(row,"principal_email","") or "").lower(),collaborator_emails=split_emails(cell(row,"collaborator_emails")))
        out.append(d)
    return out

async def preview_company_import(db,user,content):
    if user.get("role")!="admin": raise HTTPException(403,"Accès réservé aux administrateurs")
    wid=require_tenant_context(user); ws=await db.workspaces.find_one({"_id":wid})
    if not ws: raise HTTPException(403,"Workspace introuvable")
    fid=ws.get("organization_type")=="fiduciary"; rows=parse_workbook(content,fid)
    companies=await db.companies.find({"workspace_id":wid}).to_list(None); existing={str(x.get("company_code")).lower() for x in companies if x.get("company_code")}
    users=await db.users.find({"workspace_id":wid,"status":{"$ne":"inactive"}}).to_list(None); um={str(x.get("email","")).lower():str(x.get("_id") or x.get("id")) for x in users}; seen=set(); result=[]
    for r in rows:
        err=[]; warn=[]; code=(r.get("company_code") or "").lower()
        if not r["name"]: err.append("Nom de société requis")
        if r["jurisdiction"] not in {"CH","CA"}: err.append("Juridiction invalide (CH ou CA)")
        if len(r["functional_currency"])!=3: err.append("Devise fonctionnelle invalide")
        if r["company_type"] not in {"operating","holding","real_estate","nonprofit","other"}: err.append("Type de société invalide")
        if code and code in existing: err.append("Code société déjà utilisé")
        if code and code in seen: err.append("Code société dupliqué dans le fichier")
        if code: seen.add(code)
        else: warn.append("Code société non renseigné")
        pid=None; cids=[]
        if fid:
            if not r.get("mandate_code"): warn.append("Mandat non configuré: code mandat manquant")
            pe=r.get("principal_email")
            if not pe: warn.append("Mandat non configuré: responsable principal manquant")
            elif pe not in um: err.append("Responsable principal introuvable dans le workspace")
            else: pid=um[pe]
            for e in r.get("collaborator_emails",[]):
                if e not in um: err.append(f"Collaborateur introuvable: {e}")
                elif e!=pe: cids.append(um[e])
        result.append({**r,"principal_user_id":pid,"collaborator_user_ids":cids,"errors":err,"warnings":warn,"valid":not err})
    return {"total":len(result),"valid":sum(x["valid"] for x in result),"warnings":sum(bool(x["warnings"]) for x in result),"errors":sum(bool(x["errors"]) for x in result),"fiduciary":fid,"rows":result}

async def commit_company_import(db,user,content):
    p=await preview_company_import(db,user,content)
    if p["errors"]: raise HTTPException(422,detail={"message":"Import refusé: corrigez les erreurs avant validation","preview":p})
    items=[]
    for r in p["rows"]:
        keys=("name","legal_name","company_code","jurisdiction","region","functional_currency","industry","company_type","fiscal_year_start"); c=await create_company_for_admin(db,user,CompanyCreate(**{k:r.get(k) for k in keys})); m=None
        if p["fiduciary"] and r.get("mandate_code") and r.get("principal_user_id"): m=await create_mandate_for_admin(db,user,MandateCreate(company_id=c["id"],mandate_code=r["mandate_code"],principal_user_id=r["principal_user_id"],collaborator_user_ids=r["collaborator_user_ids"]))
        items.append({"company":c,"mandate":m})
    return {"created":len(items),"items":items}
